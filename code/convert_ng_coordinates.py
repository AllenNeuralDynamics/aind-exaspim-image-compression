"""
Convert Neuroglancer (x, y, z) coordinates in the evaluation dataset to
voxel (z, y, x) indices that can be used to index into the source OME-Zarr.

Reads:  /data/beta2_evaluation_dataset.xlsx
Writes: /data/beta2_evaluation_dataset_voxels.xlsx
"""

import ast
import sys

import openpyxl

sys.path.insert(0, "/root/capsule/src")
from aind_exaspim_image_compression.utils import img_util


INPUT_PATH = "/data/beta2_evaluation_dataset.xlsx"
OUTPUT_PATH = "/data/beta2_evaluation_dataset_voxels.xlsx"


def parse_coordinate(coord_str):
    """Parse a coordinate string like '(80887, -1035, 11243)' → tuple."""
    return tuple(ast.literal_eval(coord_str.strip()))


def level_path(img_path):
    """Append level '0' to a zarr path that may have a trailing slash."""
    return img_path.rstrip("/") + "/0"


def main():
    wb = openpyxl.load_workbook(INPUT_PATH)
    ws = wb.active

    # Read header and add new columns
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    voxel_col = len(headers) + 1
    ws.cell(row=1, column=voxel_col, value="Voxel Coordinate (z, y, x)")

    rows = list(ws.iter_rows(min_row=2))
    for i, row in enumerate(rows, start=2):
        brain_id = row[headers.index("Brain ID")].value
        coord_str = row[headers.index("NeuroGlancer Coordinate")].value
        img_path = row[headers.index("Image Path")].value

        lpath = level_path(img_path)
        xyz = parse_coordinate(coord_str)

        print(f"[{i-1}/{len(rows)}] Brain {brain_id} | coord {xyz} | {lpath}")
        try:
            transform = img_util.get_ome_zarr_level_transform(lpath)
            voxel_zyx = img_util.ome_zarr_coordinate_to_voxel(xyz, transform)
            print(f"         → voxel (z, y, x): {voxel_zyx}")
            ws.cell(row=i, column=voxel_col, value=str(voxel_zyx))
        except Exception as e:
            print(f"         ERROR: {e}")
            ws.cell(row=i, column=voxel_col, value=f"ERROR: {e}")

    wb.save(OUTPUT_PATH)
    print(f"\nSaved to {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
